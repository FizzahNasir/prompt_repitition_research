"""
translate_benchmarks.py
Machine-translate the SAME English benchmark items used by the Urdu/Punjabi files
into Arabic (ar), Persian (fa) and Sindhi (sd) with Google Translate, so that all
six languages share parallel items.

Item sets (aligned with datasets/urdu + datasets/punjabi):
  ARC            300  arc_challenge.csv rows 0..301 minus rows 121, 123 (exactly the
                      300 items in ARC_Urdu_GoogleTranslate.xlsx / ARC_Punjabi.xlsx,
                      same order, so '#' n here == '#' n there)
  CommonSenseQA  300  commonsenseqa.csv rows 0..299 ('#' = row+1, same as the Urdu
                      file). The Urdu/Punjabi files left 12 of these rows blank (a
                      choices-parsing artefact); they are translated here but flagged
                      'In PaUr Set' = 'No'.
  OpenBookQA     500  openbookqa.csv (full test split, same order as
                      openbookqa_urdu_test.csv). 'In Punjabi Set' flags the 495 that
                      OpenBookQA_Punjabi.xlsx contains.
  MGSM           250  mgsm_english_250.csv (all rows, same order as Urdu/Punjabi).

Output: datasets/{arabic,persian,sindhi}/{ARC,OpenBookQA,CommonSenseQA,MGSM}_<Lang>_GoogleTranslate.xlsx
using the column layout of the Urdu GoogleTranslate files (Urdu -> <Lang>), with extra
columns appended AFTER 'Status' (so positional loaders that end in `*_` still work).
OpenBookQA has no Urdu xlsx; it uses the ARC layout (same as OpenBookQA_Punjabi.xlsx).

Translation backend: deep_translator.GoogleTranslator. If that endpoint is blocked
(translate.google.com/m answers with a captcha / 429), the script falls back to the
public Google Translate "gtx" endpoint (translate.googleapis.com), i.e. the same
Google engine. Strings are batched (newline-joined, line count verified, one-by-one
fallback on mismatch) and cached on disk (JSON keyed by lang + English text), so a
crash/restart resumes without re-translating.

Usage:
  pip install deep-translator
  python translate_benchmarks.py                      # all langs, all tasks
  python translate_benchmarks.py --langs sd --tasks MGSM ARC
  python translate_benchmarks.py --cache path/to/cache.json --delay 4   # slower if Google blocks
"""

import argparse
import json
import os
import re
import sys
import tempfile
import time
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).parent
EN = ROOT / "datasets" / "english"
LANGS = {"ar": ("Arabic", "arabic"), "fa": ("Persian", "persian"), "sd": ("Sindhi", "sindhi")}
TASKS = ["ARC", "OpenBookQA", "CommonSenseQA", "MGSM"]
STATUS = "Machine translated (Google), needs review"
NLLB_STATUS = "Machine translated (NLLB-200), needs review"
NLLB_CODES = {"ar": "arb_Arab", "fa": "pes_Arab", "sd": "snd_Arab"}
DEFAULT_CACHE = Path(tempfile.gettempdir()) / "translate_benchmarks_cache.json"

ARC_SKIP = {121, 123}          # rows absent from the Urdu/Punjabi ARC files
ARC_N, CSQA_N = 300, 300

# Sentence boundary for NLLB: . ? ! then whitespace and a capital, digit or opening
# quote/bracket. Decimals (2.50) and lower-case continuations ("e.g. the") stay intact.
_SENT_SPLIT = re.compile(r"(?<=[.?!])\s+(?=[\"'(\[$]?[A-Z0-9])")
_FA_FOLD = str.maketrans("يكى", "یکی")   # NLLB's pes_Arab often emits Arabic yeh/kaf


_ABBREV = ("Mr.", "Mrs.", "Ms.", "Dr.", "St.", "Mt.", "Jr.", "Sr.", "vs.", "No.", "U.S.")


def _sentences(line: str) -> list:
    out = []
    for s in _SENT_SPLIT.split(line.strip()):
        if out and out[-1].endswith(_ABBREV):     # "Mount St. Helens", "Dr. Wertz"
            out[-1] += " " + s
        elif s:
            out.append(s)
    return out


def _degenerate(src: str, out: str) -> bool:
    """NLLB repetition loop: output far longer than the source, or a token 4+ times in a row."""
    toks = out.split()
    return (len(out) > 2.5 * len(src) + 20
            or any(toks[i] == toks[i + 1] == toks[i + 2] == toks[i + 3] for i in range(len(toks) - 3)))


# ── Source loading ────────────────────────────────────────────────────────────

def _norm(s) -> str:
    return " ".join(str(s).split()).strip().lower()


def _parse_choices(value: str) -> dict:
    """Stringified HF choices dict -> {label: text}. Evaluated with no builtins."""
    parsed = eval(str(value), {"__builtins__": {}},
                  {"array": lambda x, dtype=None: list(x), "object": object})
    return {str(k): " ".join(str(t).split()) for k, t in zip(parsed["label"], parsed["text"])}


def _mcq_items(df, qcol, rows, n_opts):
    items = []
    for r in rows:
        row = df.iloc[r]
        ch = _parse_choices(row["choices"])
        texts = list(ch.values())
        if len(texts) != n_opts:
            raise ValueError(f"row {r}: expected {n_opts} options, got {len(texts)}")
        items.append({"src_index": r, "src_id": str(row["id"]),
                      "question": " ".join(str(row[qcol]).split()),
                      "options": texts, "answer": str(row["answerKey"]).strip()})
    return items


def _pa_ur_question_set(path: Path) -> set:
    if not path.exists():
        return set()
    d = pd.read_excel(path)
    return {_norm(q) for q in d["English Question"].dropna()}


def load_source(task: str) -> list:
    if task == "ARC":
        df = pd.read_csv(EN / "arc_challenge.csv")
        rows = [r for r in range(ARC_N + len(ARC_SKIP)) if r not in ARC_SKIP]
        items = _mcq_items(df, "question", rows, 4)
        ur = _pa_ur_question_set(ROOT / "datasets/urdu/ARC_Urdu_GoogleTranslate.xlsx")
        if ur and any(_norm(it["question"]) not in ur for it in items):
            print("  WARNING: ARC item set differs from the Urdu file")
        return items
    if task == "CommonSenseQA":
        df = pd.read_csv(EN / "commonsenseqa.csv")
        items = _mcq_items(df, "question", range(CSQA_N), 5)
        ur = _pa_ur_question_set(ROOT / "datasets/urdu/CommonSenseQA_Urdu_GoogleTranslate.xlsx")
        for it in items:
            it["in_set"] = "Yes" if (not ur or _norm(it["question"]) in ur) else "No"
        return items
    if task == "OpenBookQA":
        df = pd.read_csv(EN / "openbookqa.csv")
        items = _mcq_items(df, "question_stem", range(len(df)), 4)
        pa = _pa_ur_question_set(ROOT / "datasets/punjabi/OpenBookQA_Punjabi.xlsx")
        for it in items:
            it["in_set"] = "Yes" if (not pa or _norm(it["question"]) in pa) else "No"
        return items
    if task == "MGSM":
        df = pd.read_csv(EN / "mgsm_english_250.csv")
        return [{"src_index": r, "question": " ".join(str(df.iloc[r]["question"]).split()),
                 "answer": df.iloc[r]["answer_number"],
                 "equation": df.iloc[r]["equation_solution"]} for r in range(len(df))]
    raise ValueError(task)


# ── Translation with cache ────────────────────────────────────────────────────

class Translator:
    GTX_URL = "https://translate.googleapis.com/translate_a/single"

    def __init__(self, cache_path: Path, delay: float = 2.0, max_chars: int = 4500,
                 max_lines: int = 60, backend: str = "auto", nllb_model: str = None):
        self.cache_path = Path(cache_path)
        self.delay, self.max_chars, self.max_lines = delay, max_chars, max_lines
        self.backend = backend            # auto | deep | gtx | nllb
        self.nllb_model = nllb_model
        self._nllb = None
        self.cache = {}
        if self.cache_path.exists():
            self.cache = json.loads(self.cache_path.read_text(encoding="utf-8"))
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "Mozilla/5.0"
        self.n_requests = 0
        self._dirty = 0

    def key(self, lang, text):
        # NLLB output is cached separately so it never mixes with Google translations
        # (nllb2: sentence-level; nllb- entries were whole-line and dropped sentences)
        return f"nllb2-{lang}\t{text}" if self.backend == "nllb" else f"{lang}\t{text}"

    # -- NLLB-200 (local model, no network quota) --
    def _nllb_translate(self, texts, lang, batch_size=16, strict=False):
        if self._nllb is None:
            import torch
            from transformers import AutoModelForSeq2SeqLM, AutoTokenizer
            tok = AutoTokenizer.from_pretrained(self.nllb_model, src_lang="eng_Latn")
            model = AutoModelForSeq2SeqLM.from_pretrained(self.nllb_model)
            model.to("cuda" if torch.cuda.is_available() else "cpu").eval()
            self._nllb = (tok, model, torch)
            print(f"  NLLB model loaded: {self.nllb_model} on {model.device}", flush=True)
        tok, model, torch = self._nllb
        target = tok.convert_tokens_to_ids(NLLB_CODES[lang])
        # Retries of looping outputs get a repetition penalty; normal outputs don't, since
        # blocking repeated n-grams would also mangle legitimately repeated names/numbers.
        extra = {"repetition_penalty": 1.3, "no_repeat_ngram_size": 3} if strict else {}
        out = []
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            enc = tok(batch, return_tensors="pt", padding=True, truncation=True,
                      max_length=512).to(model.device)
            with torch.no_grad():
                gen = model.generate(**enc, forced_bos_token_id=target, num_beams=4,
                                     max_new_tokens=int(enc["input_ids"].shape[1] * 2) + 20,
                                     **extra)
            out.extend(tok.batch_decode(gen, skip_special_tokens=True))
        if lang == "fa":
            out = [o.translate(_FA_FOLD) for o in out]
        bad = [k for k, (s, o) in enumerate(zip(texts, out)) if _degenerate(s, o)]
        if bad and not strict:
            print(f"    {len(bad)} looping outputs, retrying with a repetition penalty", flush=True)
            for k, o in zip(bad, self._nllb_translate([texts[k] for k in bad], lang, strict=True)):
                out[k] = o
                if _degenerate(texts[k], o):
                    print(f"    WARNING still looping: {texts[k][:80]!r}", flush=True)
        return out

    def _translate_nllb(self, todo, lang):
        # NLLB translates one sentence well but drops sentences from longer input, so
        # translate each sentence of each line separately and re-join.
        lines = list(dict.fromkeys(l for t in todo for l in t.split("\n") if l.strip()))
        sents = list(dict.fromkeys(s for l in lines for s in _sentences(l)))
        done, step = {}, 128
        for i in range(0, len(sents), step):
            for src, out in zip(sents[i:i + step], self._nllb_translate(sents[i:i + step], lang)):
                done[src] = out
            for t in todo:
                if self.key(lang, t) in self.cache:
                    continue
                if all(s in done for l in t.split("\n") for s in _sentences(l)):
                    self._store(lang, t, "\n".join(" ".join(done[s] for s in _sentences(l)) if l.strip() else l
                                                   for l in t.split("\n")))
            self.save()
            print(f"    {min(i + step, len(sents))}/{len(sents)} sentences", flush=True)

    def save(self):
        tmp = self.cache_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.cache, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, self.cache_path)
        self._dirty = 0

    # -- raw backends --
    def _deep(self, text, lang):
        from deep_translator import GoogleTranslator
        return GoogleTranslator(source="en", target=lang).translate(text)

    def _gtx(self, text, lang):
        r = self.session.post(self.GTX_URL, params={"client": "gtx", "sl": "en", "tl": lang, "dt": "t"},
                              data={"q": text}, timeout=60)
        if r.status_code == 429:
            raise RuntimeError("429 Too Many Requests")
        r.raise_for_status()
        return "".join(seg[0] for seg in r.json()[0] if seg and seg[0])

    def _raw(self, text, lang):
        """One request, with retries + exponential backoff."""
        wait = 120.0
        for attempt in range(40):            # ~9 h of patience before giving up
            time.sleep(self.delay)
            self.n_requests += 1
            try:
                if self.backend in ("auto", "deep"):
                    try:
                        out = self._deep(text, lang)
                    except Exception as e:  # blocked / captcha / 429 on translate.google.com/m
                        if self.backend == "deep":
                            raise
                        print(f"    deep_translator unavailable ({type(e).__name__}: {str(e)[:80]}); "
                              f"switching to Google gtx endpoint")
                        self.backend = "gtx"
                        out = self._gtx(text, lang)
                else:
                    out = self._gtx(text, lang)
                if out is None:
                    raise RuntimeError("empty response")
                return out
            except Exception as e:
                print(f"    request failed ({type(e).__name__}: {str(e)[:100]}); "
                      f"retry {attempt + 1} in {wait:.0f}s", flush=True)
                time.sleep(wait)
                wait = min(wait * 1.5, 900)
        raise RuntimeError("giving up after repeated failures; rerun to resume from cache")

    def _store(self, lang, src, out):
        self.cache[self.key(lang, src)] = out.strip()
        self._dirty += 1

    def _translate_chunk(self, chunk, lang):
        if len(chunk) == 1:
            self._store(lang, chunk[0], self._raw(chunk[0], lang))
            return
        out = self._raw("\n".join(chunk), lang)
        parts = [p for p in out.split("\n")]
        parts = [p for p in parts if p.strip()] if len(parts) != len(chunk) else parts
        if len(parts) == len(chunk) and all(p.strip() for p in parts):
            for s, p in zip(chunk, parts):
                self._store(lang, s, p)
        else:
            print(f"    batch line mismatch ({len(parts)} vs {len(chunk)}); translating one-by-one")
            for s in chunk:
                self._store(lang, s, self._raw(s, lang))

    def translate_many(self, texts, lang, label=""):
        todo = list(dict.fromkeys(t for t in texts if t and self.key(lang, t) not in self.cache))
        if todo:
            print(f"  [{lang}] {label}: {len(todo)} new strings to translate "
                  f"({len(set(texts)) - len(todo)} cached)")
        if self.backend == "nllb":
            if todo:
                self._translate_nllb(todo, lang)
            return {t: self.cache.get(self.key(lang, t), "") for t in texts}
        chunk, size, done = [], 0, 0
        for t in todo + [None]:
            flush = t is None or (chunk and (size + len(t) > self.max_chars or len(chunk) >= self.max_lines))
            if flush and chunk:
                self._translate_chunk(chunk, lang)
                done += len(chunk)
                self.save()
                if done % 400 < len(chunk):
                    print(f"    {done}/{len(todo)}")
                chunk, size = [], 0
            if t is not None:
                if "\n" in t:            # never batch multi-line strings
                    self._store(lang, t, self._raw(t, lang))
                    continue
                chunk.append(t)
                size += len(t) + 1
        return {t: self.cache.get(self.key(lang, t), "") for t in texts}

    def retranslate_single(self, text, lang):
        out = self._raw(text, lang).strip()
        return out


# ── Digits / number check (MGSM only) ─────────────────────────────────────────

_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")


def normalize_digits(s: str) -> str:
    s = s.translate(_DIGITS)
    s = re.sub(r"(?<=\d)٬(?=\d)", ",", s)    # Arabic thousands separator
    s = re.sub(r"(?<=\d)٫(?=\d)", ".", s)    # Arabic decimal separator
    return s


def numbers_in(s: str) -> set:
    out = set()
    for m in re.findall(r"(?:\d+(?:[,،]\d{3})*)?(?:\.\d+)?", s):
        if not m or not any(ch.isdigit() for ch in m):
            continue
        v = re.sub(r"[,،]", "", m)
        v = "0" + v if v.startswith(".") else v
        if "." in v:
            v = v.rstrip("0").rstrip(".")
        out.add(v)
    return out


def missing_numbers(en: str, tr: str) -> list:
    return sorted(numbers_in(en) - numbers_in(normalize_digits(tr)), key=lambda x: float(x))


# Number words Google tends to substitute for digits (Arabic duals "ساعتين", "80 ألف", ...).
_NUM_WORDS = {
    "ar": {1: "واحد|واحدة", 2: r"اثن|ين\b|ان\b", 3: "ثلاث", 4: "أربع|اربع", 5: "خمس", 6: "ست",
           7: "سبع", 8: "ثمان", 9: "تسع", 10: "عشر", 12: "اثني عشر|اثنا عشر|دزينة", 20: "عشرين|عشرون",
           30: "ثلاثين|ثلاثون", 40: "أربعين|أربعون", 50: "خمسين|خمسون", 60: "ستين|ستون", 70: "سبعين",
           75: "خمسة وسبعين|خمس وسبعين", 80: "ثمانين", 90: "تسعين", 100: "مائة|مئة", 1000: "ألف"},
    "fa": {1: "یک", 2: "دو", 3: "سه", 4: "چهار", 5: "پنج", 6: "شش", 7: "هفت", 8: "هشت", 9: "نه",
           10: "ده", 12: "دوازده", 20: "بیست", 30: "سی", 40: "چهل", 50: "پنجاه", 60: "شصت", 100: "صد",
           1000: "هزار"},
    "sd": {1: "هڪ", 2: "ٻه|ٻن|ٻئي", 3: "ٽي|ٽن", 4: "چار", 5: "پنج", 6: "ڇهه|ڇه", 7: "ست", 8: "اٺ",
           9: "نو", 10: "ڏهه|ڏه", 12: "ٻارهن|درجن", 20: "ويه", 30: "ٽيه", 40: "چاليه", 50: "پنجاهه",
           100: "سو", 1000: "هزار"},
}


def classify_missing(miss: list, tr: str, lang: str, en: str = ""):
    """Split missing digits into (word_form, truly_missing). A number counts as 'word_form'
    when a number word for it (or 'N thousand' for multiples of 1000) occurs in the text."""
    words, t = _NUM_WORDS.get(lang, {}), normalize_digits(tr)
    word_form, missing = [], []
    in_fraction = {x for f in re.findall(r"\d+/\d+", en) for x in f.split("/")}
    for m in miss:
        v = float(m)
        ok = False
        if m in in_fraction and not re.search(r"\d+/\d+", t):   # "1/4" -> "a quarter"
            ok = True
        elif v.is_integer() and int(v) in words and re.search(words[int(v)], t):
            ok = True
        elif v.is_integer() and v >= 1000 and v % 1000 == 0 and 1000 in words:
            k = str(int(v // 1000))
            ok = bool(re.search(rf"(?<!\d){k}\s*(?:{words[1000]})", t))
        (word_form if ok else missing).append(m)
    return word_form, missing


# ── QC helpers ────────────────────────────────────────────────────────────────

_ARABIC_SCRIPT = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
SINDHI_CORE = set("ڪڏٻڄڃ")
SINDHI_ALL = set("ڪڏٻڄڃٽڊڙڳڱڦڀٿڌڇڻ")


def has_arabic(s): return bool(_ARABIC_SCRIPT.search(s or ""))


# ── Build one file ────────────────────────────────────────────────────────────

def write_xlsx(path: Path, sheet: str, header: list, rows: list):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, start=1):
        col = ws.cell(1, i).column_letter
        ws.column_dimensions[col].width = 50 if "Question" in h else (22 if len(h) <= 8 else 16)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=False, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build(task: str, lang: str, tr: Translator) -> dict:
    name, folder = LANGS[lang]
    items = load_source(task)
    nllb = tr.backend == "nllb"
    engine, status = ("NLLB", NLLB_STATUS) if nllb else ("Google Translate", STATUS)
    out_path = ROOT / "datasets" / folder / f"{task}_{name}_{'NLLB' if nllb else 'GoogleTranslate'}.xlsx"
    qc = {"file": str(out_path.relative_to(ROOT)), "rows": len(items)}

    if task == "MGSM":
        tmap = tr.translate_many([it["question"] for it in items], lang, f"{task}")
        header = ["#", "English Question", f"{engine} {name}", "Answer", "Equation", "Status",
                  "Source Index", "Number Check"]
        rows, mism, wordonly = [], [], 0
        for n, it in enumerate(items, start=1):
            t = normalize_digits(tmap[it["question"]])
            miss = missing_numbers(it["question"], t)
            word_form, hard = classify_missing(miss, t, lang, it["question"])
            if miss:
                mism.append((n, word_form, hard))
                wordonly += not hard
            check = "OK"
            if miss:
                parts = []
                if hard:
                    parts.append("MISSING: " + ", ".join(hard))
                if word_form:
                    parts.append("AS WORD: " + ", ".join(word_form))
                check = "; ".join(parts)
            ans = it["answer"]
            ans = int(ans) if float(ans).is_integer() else float(ans)
            eq = None if pd.isna(it["equation"]) else it["equation"]
            rows.append([n, it["question"], t, ans, eq, status, it["src_index"], check])
        qc.update(empty=sum(1 for r in rows if not r[2].strip()),
                  number_mismatch=len(mism), number_mismatch_rows=mism, number_word_only=wordonly,
                  arabic_script=sum(has_arabic(r[2]) for r in rows),
                  native_texts=[r[2] for r in rows])
        samples = [(r[1], r[2], r[3]) for r in rows[:2]]
    else:
        letters = "ABCDE"[:len(items[0]["options"])]
        strings = []
        for it in items:
            strings.append(it["question"])
            strings.extend(it["options"])
        tmap = tr.translate_many(strings, lang, task)
        header = (["#", "English Question", f"{name} Question"] + [f"Eng {l}" for l in letters]
                  + [f"{name} {l}" for l in letters] + ["Answer", "Status", "Source ID", "Source Index"]
                  + (["In PaUr Set"] if task == "CommonSenseQA" else [])
                  + (["In Punjabi Set"] if task == "OpenBookQA" else []) + ["Duplicate Options"])
        rows, dup, native_texts, empty = [], 0, [], 0
        for n, it in enumerate(items, start=1):
            q = tmap[it["question"]]
            opts = [tmap[o] for o in it["options"]]
            empty += sum(1 for s in [q] + opts if not s.strip())
            is_dup = len({o.strip() for o in opts}) < len(opts)
            en_dup = len({o.strip().lower() for o in it["options"]}) < len(opts)
            dup += is_dup
            flag = ("Yes (also in English)" if en_dup else "Yes") if is_dup else ""
            extra = [it["in_set"]] if task in ("CommonSenseQA", "OpenBookQA") else []
            rows.append([n, it["question"], q] + it["options"] + opts
                        + [it["answer"], status, it["src_id"], it["src_index"]] + extra + [flag])
            native_texts.append(" ".join([q] + opts))
        qc.update(empty=empty, duplicate_option_items=dup,
                  arabic_script=sum(has_arabic(t) for t in native_texts), native_texts=native_texts)
        if task in ("CommonSenseQA", "OpenBookQA"):
            qc["in_set_yes"] = sum(it["in_set"] == "Yes" for it in items)
        samples = [(r[1], r[2], dict(zip(letters, r[3 + len(letters):3 + 2 * len(letters)])), r[3 + 2 * len(letters)])
                   for r in rows[:2]]

    write_xlsx(out_path, f"{task} {name}", header, rows)
    qc["columns"] = header
    qc["samples"] = samples
    return qc


def report(qc: dict, lang: str):
    texts = qc.pop("native_texts")
    n = len(texts)
    print(f"\n== {qc['file']}  rows={qc['rows']}")
    print(f"   empty translations: {qc['empty']}   Arabic-script items: {qc['arabic_script']}/{n}")
    if "duplicate_option_items" in qc:
        print(f"   items with identical translated options: {qc['duplicate_option_items']}")
    if "in_set_yes" in qc:
        print(f"   items also in pa/ur set: {qc['in_set_yes']}/{n}")
    if "number_mismatch" in qc:
        print(f"   MGSM items whose English digits are not all present as digits: {qc['number_mismatch']} "
              f"(only number-word/dual forms: {qc['number_word_only']}; "
              f"needs review: {qc['number_mismatch'] - qc['number_word_only']})")
        for row_n, word_form, hard in qc["number_mismatch_rows"]:
            if hard:
                print(f"     #{row_n}: missing {hard}" + (f" (as word: {word_form})" if word_form else ""))
    if lang == "sd":
        core = sum(any(ch in SINDHI_CORE for ch in t) for t in texts)
        anyl = sum(any(ch in SINDHI_ALL for ch in t) for t in texts)
        qc["sindhi_core_share"] = core / n
        qc["sindhi_any_share"] = anyl / n
        print(f"   Sindhi letters (ڪ ڏ ٻ ڄ ڃ) present: {core}/{n} = {core / n:.1%}; "
              f"any Sindhi-specific letter: {anyl}/{n} = {anyl / n:.1%}")
    for s in qc.pop("samples"):
        print("   sample EN:", s[0][:160])
        print("          TR:", s[1][:200])
        if len(s) == 4:
            print("          opts:", s[2], " answer:", s[3])
        else:
            print("          answer:", s[2])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--langs", nargs="+", default=list(LANGS), choices=list(LANGS))
    ap.add_argument("--tasks", nargs="+", default=TASKS, choices=TASKS)
    ap.add_argument("--cache", default=str(DEFAULT_CACHE))
    ap.add_argument("--delay", type=float, default=2.0, help="seconds between requests")
    ap.add_argument("--backend", default="auto", choices=["auto", "deep", "gtx", "nllb"])
    ap.add_argument("--nllb-model", default="facebook/nllb-200-distilled-1.3B",
                    help="NLLB checkpoint for --backend nllb (writes *_NLLB.xlsx)")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

    t0 = time.time()
    tr = Translator(Path(args.cache), delay=args.delay, backend=args.backend,
                    nllb_model=args.nllb_model)
    print(f"cache: {args.cache} ({len(tr.cache)} entries)")
    summary = []
    try:
        for lang in args.langs:
            for task in args.tasks:
                qc = build(task, lang, tr)
                report(qc, lang)
                summary.append(qc)
    finally:
        tr.save()
    print(f"\nDone in {time.time() - t0:.0f}s, {tr.n_requests} requests, backend={tr.backend}")


if __name__ == "__main__":
    main()
